from pathlib import Path
from graphify.detect import classify_file, count_words, detect, FileType, _looks_like_paper

FIXTURES = Path(__file__).parent / "fixtures"

def test_classify_python():
    assert classify_file(Path("foo.py")) == FileType.CODE

def test_classify_typescript():
    assert classify_file(Path("bar.ts")) == FileType.CODE

def test_classify_markdown():
    assert classify_file(Path("README.md")) == FileType.DOCUMENT

def test_classify_pdf():
    assert classify_file(Path("paper.pdf")) == FileType.PAPER

def test_classify_unknown_returns_none():
    assert classify_file(Path("archive.zip")) is None

def test_classify_image():
    assert classify_file(Path("screenshot.png")) == FileType.IMAGE
    assert classify_file(Path("design.jpg")) == FileType.IMAGE
    assert classify_file(Path("diagram.webp")) == FileType.IMAGE

def test_count_words_sample_md():
    words = count_words(FIXTURES / "sample.md")
    assert words > 5

def test_detect_finds_fixtures():
    result = detect(FIXTURES)
    assert result["total_files"] >= 2
    assert "code" in result["files"]
    assert "document" in result["files"]

def test_detect_warns_small_corpus():
    result = detect(FIXTURES)
    assert result["needs_graph"] is False
    assert result["warning"] is not None

def test_detect_skips_dotfiles():
    result = detect(FIXTURES)
    for files in result["files"].values():
        for f in files:
            assert "/." not in f


def test_classify_md_paper_by_signals(tmp_path):
    """A .md file with enough paper signals should classify as PAPER."""
    paper = tmp_path / "paper.md"
    paper.write_text(
        "# Abstract\n\nWe propose a new method. See [1] and [23].\n"
        "This work was published in the Journal of AI. ArXiv preprint.\n"
        "See Equation 3 for details. \\cite{vaswani2017}.\n"
    )
    assert classify_file(paper) == FileType.PAPER


def test_classify_md_doc_without_signals(tmp_path):
    """A plain .md file without paper signals should stay DOCUMENT."""
    doc = tmp_path / "notes.md"
    doc.write_text("# My Notes\n\nHere are some notes about the project.\n")
    assert classify_file(doc) == FileType.DOCUMENT


def test_classify_attention_paper():
    """The real attention paper file should be classified as PAPER."""
    paper_path = Path("/home/safi/graphify_eval/papers/attention_is_all_you_need.md")
    if paper_path.exists():
        result = classify_file(paper_path)
        assert result == FileType.PAPER

def test_classify_docx():
    assert classify_file(Path("report.docx")) == FileType.DOCUMENT

def test_classify_xlsx():
    assert classify_file(Path("data.xlsx")) == FileType.DOCUMENT

def test_classify_pptx():
    assert classify_file(Path("slides.pptx")) == FileType.DOCUMENT

import subprocess
import pytest

def test_extract_office_text_docx(tmp_path):
    from graphify.detect import extract_office_text
    result = subprocess.run(["pandoc", "--version"], capture_output=True)
    if result.returncode != 0:
        pytest.skip("pandoc not available")
    src = tmp_path / "input.md"
    src.write_text("# Title\n\nHello world content here.")
    docx = tmp_path / "test.docx"
    subprocess.run(["pandoc", str(src), "-o", str(docx)], check=True)
    text = extract_office_text(docx)
    assert "Hello" in text or "Title" in text

def test_extract_office_text_xlsx(tmp_path):
    pytest.importorskip("openpyxl")
    import openpyxl
    from graphify.detect import extract_office_text
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Revenue"
    ws.append(["Month", "Amount", "Region"])
    ws.append(["Jan", 1000, "North"])
    path = tmp_path / "data.xlsx"
    wb.save(path)
    text = extract_office_text(path)
    assert "Revenue" in text
    assert "Month" in text
    assert "Amount" in text

def test_extract_office_text_pptx(tmp_path):
    pytest.importorskip("pptx")
    from pptx import Presentation
    from graphify.detect import extract_office_text
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Q4 Strategy"
    path = tmp_path / "deck.pptx"
    prs.save(path)
    text = extract_office_text(path)
    assert "Q4 Strategy" in text

def test_extract_office_text_returns_empty_on_error(tmp_path):
    from graphify.detect import extract_office_text
    fake = tmp_path / "corrupt.docx"
    fake.write_bytes(b"not a real docx")
    result = extract_office_text(fake)
    assert isinstance(result, str)

def test_count_words_docx(tmp_path):
    result = subprocess.run(["pandoc", "--version"], capture_output=True)
    if result.returncode != 0:
        pytest.skip("pandoc not available")
    src = tmp_path / "input.md"
    src.write_text("# Title\n\nThis document has several words in it.")
    docx = tmp_path / "test.docx"
    subprocess.run(["pandoc", str(src), "-o", str(docx)], check=True)
    from graphify.detect import count_words
    assert count_words(docx) > 5
